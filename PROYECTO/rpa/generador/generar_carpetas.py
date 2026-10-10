"""Genera una carpeta de prueba para el robot de UiPath, parecida a la que se entrega en la calificación.

Imita el enunciado y los archivos de ejemplo del foro (PROYECTO/rpa/ejemplos/):
- Carpetas "clientes - ...", "proveedores - ...", "reclamos - ...", "registro - ...", "productos - ..." con
  subcarpetas, y archivos .xlsx y .xls con nombres poco prácticos.
- Hojas "clientes", "proveedor", "reclamos", "registros", "productos" y otras. Solo importan las hojas
  "clientes" y "productos", estén en la carpeta que estén.
- Mismos encabezados y formatos que los ejemplos: "Name*", "Company Type*" con Company/Person, país con
  código ISO, código postal numérico, código de barras numérico, "True"/"False" en "Está publicado",
  "Related Company" y "Product Values" siempre vacías.
- Trampas: carpeta "clientes - ..." sin hoja clientes, hoja productos dentro de "proveedores - ...",
  duplicados entre archivos, filas sin campos obligatorios, filas vacías con formato, servicio con
  cantidad a la mano, un .txt y un archivo temporal "~$" de Excel.

Al final imprime cuántos clientes y productos únicos y válidos debe cargar el robot.

Requisitos: pip install openpyxl xlwt
Uso: python generar_carpetas.py [carpeta_destino]   (por defecto PROYECTO/rpa/carpeta_prueba)
"""
import shutil
import sys
from pathlib import Path

import openpyxl
import xlwt
from openpyxl.styles import Font, PatternFill

DESTINO = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "carpeta_prueba"

ENC_CLIENTES = ["Name*", "Company Type*", "Related Company", "Email", "Phone", "Street", "Street2", "City",
                "State", "Zip", "Country", "Tax ID", "Website", "Tags", "Reference", "Notes"]
ENC_PRODUCTOS = ["External ID", "Name", "Product Type", "Internal Reference", "Barcode", "Sales Price", "Cost",
                 "Weight", "Sales Description", "Product Values", "Cantidad a la mano", "Está publicado"]


def cliente(nombre, tipo, email=None, telefono=None, calle=None, ciudad=None, estado=None, zip_=None, pais=None,
            nit=None, web=None, etiquetas=None, ref=None, notas=None):
    # "Related Company" siempre vacía (foro).
    return [nombre, tipo, None, email, telefono, calle, None, ciudad, estado, zip_, pais, nit, web, etiquetas, ref, notas]


def producto(ext_id, nombre, tipo, ref=None, barcode=None, precio=None, costo=None, peso=None, desc=None,
             cantidad=None, publicado=None):
    # "Product Values" siempre vacía (foro).
    return [ext_id, nombre, tipo, ref, barcode, precio, costo, peso, desc, None, cantidad, publicado]


# --- Datos ---------------------------------------------------------------------------------------
CLIENTES_A = [
    cliente("Supermercado La Esquina, S.A.", "Company", "compras@laesquina.example", "+502 2334 1100",
            "5a. Avenida 12-40, Zona 1", "Ciudad de Guatemala", None, 1001, "GT", "4567890-1",
            "https://laesquina.example", "Mayorista", "RPA-CLI-001", "Cliente <b>mayorista</b> del centro"),
    cliente("Abarrotería Doña Tere", "Company", "tere@abarroteria.example", "+502 7765 2200", "Calle Real 8",
            "Antigua Guatemala", None, 3001, "GT", None, None, "Minorista,Cliente RPA", "RPA-CLI-002"),
    cliente("Lucía Méndez Orellana", "Person", "lucia.mendez@correo.example", "+502 5512 3344",
            "Calzada Roosevelt 15-20", "Mixco", None, 1057, "GT"),
    cliente("Carlos Ajú Tzoc", "Person", "carlos.aju@correo.example", None, None, "Quetzaltenango", None, None, "GT"),
    cliente("Restaurante El Comal", "Company", "pedidos@elcomal.example", "+52 55 4455 6677", "Av. Juárez 220",
            "Ciudad de México", None, 6000, "MX", None, "https://elcomal.example", "Mayorista", "RPA-CLI-005"),
    cliente("Ana Sofía Ramírez", "Person", "anasofia@correo.example", "+52 33 1122 3344", None, "Guadalajara",
            "Jalisco", 44100, "MX"),
]
CLIENTES_B = [
    # Duplicado de un cliente de CLIENTES_A (mismo nombre y correo): se carga una sola vez.
    cliente("Abarrotería Doña Tere", "Company", "tere@abarroteria.example", "+502 7765 2200", "Calle Real 8",
            "Antigua Guatemala", None, 3001, "GT", None, None, "Minorista,Cliente RPA", "RPA-CLI-002"),
    cliente("Pupusería Los Volcanes", "Company", "ventas@losvolcanes.example", "+503 2222 3344",
            "Blvd. de Los Héroes 300", "San Salvador", None, 1101, "SV", None, None, "Cliente RPA", "RPA-CLI-007"),
    cliente("José Hernández Portillo", "Person", "jose.hp@correo.example", "+503 7788 9900", None, "Santa Ana",
            None, 2201, "SV"),
    cliente(None, "Person", "sin.nombre@correo.example"),          # sin Name*: se rechaza
    cliente("Cliente sin tipo", None, "sin.tipo@correo.example"),  # sin Company Type*: se rechaza
]
CLIENTES_C = [
    cliente("Hotel Casa del Lago", "Company", "compras@casadellago.example", "+502 7762 0000", "Km 3 Panajachel",
            "Panajachel", None, 7010, "GT", "7788990-0", "https://casadellago.example", "Mayorista", "RPA-CLI-010",
            "Pide entregas los lunes"),
    cliente("María José Cifuentes", "Person", "mjcifuentes@correo.example", None, None, None, None, None, None),
]

PRODUCTOS_A = [
    producto("RPA_PRUEBA_PROD_001", "Tortillas de maíz amarillo 30 u", "Goods", "RPA-001", 7401234500018, 12.5, 8,
             1.1, "Tortillas frescas hechas cada mañana.", 40, "True"),
    producto("RPA_PRUEBA_PROD_002", "Frijol volteado en bolsa 400 g", "Goods", "RPA-002", 7401234500025, 9.75, 6.2,
             0.4, "Frijol negro volteado, listo para servir.", 60, "True"),
    producto("RPA_PRUEBA_PROD_003", "Chocolate de mesa artesanal 4 tabletas", "Goods", "RPA-003", 7401234500032,
             22, 14, 0.25, "Chocolate de cacao guatemalteco con canela.", 25, "True"),
    producto("RPA_PRUEBA_PROD_004", "Servicio de entrega express", "Service", None, None, 35, None, None,
             "Entrega en menos de 2 horas en la ciudad.", 75, "True"),   # servicio con cantidad: no lleva existencias
    producto("RPA_PRUEBA_PROD_005", "Canasta navideña ejecutiva", "Goods", "RPA-005", None, 450, 310, 6.5,
             "Canasta con productos seleccionados.", None, "False"),
]
PRODUCTOS_B = [
    # Duplicado por External ID: se carga una sola vez.
    producto("RPA_PRUEBA_PROD_002", "Frijol volteado en bolsa 400 g", "Goods", "RPA-002", 7401234500025, 9.75, 6.2,
             0.4, "Frijol negro volteado, listo para servir.", 60, "True"),
    producto("RPA_PRUEBA_PROD_006", "Café molido de altura 454 g", "Goods", "RPA-006", 7401234500063, 48, 31, 0.454,
             "Café de Huehuetenango, tueste medio.", 30, "True"),
    producto("RPA_PRUEBA_PROD_007", "Asesoría de recetas para eventos", "Service", None, None, 150, None, None, None,
             None, None),
    producto(None, "Producto sin External ID", "Goods", "RPA-008", None, 10, 5),   # sin External ID: se rechaza
    producto("RPA_PRUEBA_PROD_009", "Producto sin tipo", None, "RPA-009", None, 10, 5),  # sin Product Type: se rechaza
]
PRODUCTOS_C = [
    producto("RPA_PRUEBA_PROD_010", "Horchata en polvo 400 g", "Goods", "RPA-010", 7401234500100, 18.5, 11, 0.4,
             "Bebida tradicional de morro y arroz.", 45, "True"),
]

OTRAS_HOJAS = {
    "proveedor": (["Proveedor", "Contacto", "Teléfono"], [["Distribuidora La Cosecha", "Mario Pérez", "+502 2200 1100"]]),
    "reclamos": (["Fecha", "Cliente", "Motivo"], [["2026-09-14", "Lucía Méndez", "Producto vencido"]]),
    "registros": (["Fecha", "Sucursal", "Ventas Q"], [["2026-10-01", "GT", 15234.5], ["2026-10-01", "SV", 8450]]),
    "resumen": (["Total", "Observación"], [[12, "Revisar con contabilidad"]]),
    "Hoja1": (["Notas"], [["Archivo de trabajo, no usar"]]),
}


# --- Escritura -----------------------------------------------------------------------------------
def escribir_xlsx(ruta, hojas, filas_vacias=0):
    """hojas: {nombre: (encabezados, filas)}. filas_vacias imita las filas con formato pero sin datos del ejemplo."""
    libro = openpyxl.Workbook()
    libro.remove(libro.active)
    for nombre, (encabezados, filas) in hojas.items():
        hoja = libro.create_sheet(nombre)
        hoja.append(encabezados)
        for celda in hoja[1]:
            celda.font = Font(bold=True)
        for fila in filas:
            hoja.append(fila)
        for n in range(filas_vacias):
            hoja.cell(row=len(filas) + 2 + n, column=1).fill = PatternFill("solid", fgColor="FFFFFF")
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta)


def escribir_xls(ruta, hojas):
    libro = xlwt.Workbook(encoding="utf-8")
    negrita = xlwt.easyxf("font: bold on")
    for nombre, (encabezados, filas) in hojas.items():
        hoja = libro.add_sheet(nombre)
        for c, valor in enumerate(encabezados):
            hoja.write(0, c, valor, negrita)
        for f, fila in enumerate(filas, start=1):
            for c, valor in enumerate(fila):
                if valor is not None:
                    hoja.write(f, c, valor)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro.save(str(ruta))


def otras(*nombres):
    return {n: OTRAS_HOJAS[n] for n in nombres}


def main():
    if DESTINO.exists():
        shutil.rmtree(DESTINO)
    raiz = DESTINO

    # (ruta relativa, hojas, formato)
    archivos = [
        ("clientes - base historica 2024 NO BORRAR/clientes - copia final(2) revisada.xlsx",
         {"clientes": (ENC_CLIENTES, CLIENTES_A), **otras("resumen")}, "xlsx"),
        ("clientes - base historica 2024 NO BORRAR/clientes - borrador sin terminar.xlsx",
         otras("registros", "Hoja1"), "xlsx"),
        ("clientes - base historica 2024 NO BORRAR/zona norte/viejos/clientes - listado norte v3 ULTIMO.xls",
         {"clientes": (ENC_CLIENTES, CLIENTES_C), **otras("registros")}, "xls"),
        ("proveedores - contactos de compras/proveedores - lista 2025.xlsx", otras("proveedor"), "xlsx"),
        ("proveedores - contactos de compras/proveedores - mixto revisar.xlsx",
         {**otras("proveedor"), "productos": (ENC_PRODUCTOS, PRODUCTOS_C)}, "xlsx"),
        ("reclamos - atencion al cliente sucursales/reclamos - septiembre.xlsx", otras("reclamos"), "xlsx"),
        ("reclamos - atencion al cliente sucursales/reclamos - con clientes nuevos.xls",
         {**otras("reclamos"), "clientes": (ENC_CLIENTES, CLIENTES_B)}, "xls"),
        ("registro - ventas diarias/registro - octubre.xlsx", otras("registros"), "xlsx"),
        ("registro - ventas diarias/registro - archivo muerto/registro - mezcla clientes y productos FINAL.xlsx",
         {"registros": OTRAS_HOJAS["registros"], "clientes": (ENC_CLIENTES, CLIENTES_B[:3]),
          "productos": (ENC_PRODUCTOS, PRODUCTOS_B)}, "xlsx"),
        ("productos - catalogo nuevo temporada/productos - lanzamientos.xls",
         {"productos": (ENC_PRODUCTOS, PRODUCTOS_A), **otras("resumen")}, "xls"),
        ("productos - catalogo nuevo temporada/productos - precios viejos NO USAR.xlsx", otras("Hoja1"), "xlsx"),
    ]
    for ruta, hojas, formato in archivos:
        destino = raiz / ruta
        if formato == "xls":
            escribir_xls(destino, hojas)
        else:
            escribir_xlsx(destino, hojas, filas_vacias=200 if "copia final" in ruta else 0)

    # Archivos que el robot debe ignorar.
    (raiz / "productos - catalogo nuevo temporada/notas - leeme.txt").write_text(
        "Los precios de este archivo están desactualizados.\n", encoding="utf-8")
    (raiz / "clientes - base historica 2024 NO BORRAR/~$clientes - copia final(2) revisada.xlsx").write_bytes(b"\x00" * 165)

    # Resultado esperado tras consolidar (únicos y válidos).
    def validos_clientes(filas):
        return {(f[0], f[3]) for f in filas if f[0] and f[1]}

    def validos_productos(filas):
        return {f[0] for f in filas if f[0] and f[1] and f[2]}

    clientes = validos_clientes(CLIENTES_A + CLIENTES_B + CLIENTES_C)
    productos = validos_productos(PRODUCTOS_A + PRODUCTOS_B + PRODUCTOS_C)
    con_existencias = {f[0] for f in PRODUCTOS_A + PRODUCTOS_B + PRODUCTOS_C
                       if f[0] in productos and f[2] == "Goods" and f[10]}
    total = sum(1 for _ in raiz.rglob("*") if _.is_file())
    print(f"Carpeta de prueba: {raiz} ({total} archivos)")
    print(f"Esperado: {len(clientes)} clientes, {len(productos)} productos, {len(con_existencias)} con existencias, "
          f"2 clientes y 2 productos rechazados por campos obligatorios")


if __name__ == "__main__":
    main()
