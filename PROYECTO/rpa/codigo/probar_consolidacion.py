"""Prueba, sin UiPath, el código C# que el robot ejecuta en sus actividades Invoke Code.

Emula los pasos del robot: recorre la carpeta, ignora los archivos temporales "~$", abre cada .xls/.xlsx,
lee solo las hojas llamadas exactamente "clientes" y "productos" con valores sin formato (como Read Range),
ejecuta AgregarOrigen.cs por hoja, une las hojas (Merge Data Table) y ejecuta Consolidar.cs.
El C# se compila con el compilador de .NET Framework que trae Windows, así que se prueba exactamente el mismo
código que se pega en UiPath. No se conecta a Odoo.

Requisitos: Windows, pip install openpyxl xlrd
Uso: python probar_consolidacion.py <carpeta_raiz> [carpeta_salida]
     (por defecto la salida va a PROYECTO/rpa/salida/prueba/: clientes, productos, existencias y rechazos en .tsv)
"""
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
import xlrd

AQUI = Path(__file__).resolve().parent
CSC = Path(r"C:\Windows\Microsoft.NET\Framework64\v4.0.30319\csc.exe")
HOJAS = ("clientes", "productos")
UBICACION = "GT/Existencias"

PLANTILLA = r"""
using System; using System.Collections.Generic; using System.Data; using System.IO; using System.Linq; using System.Text;
class Prueba {
    static void AgregarOrigen(DataTable dtHoja, string archivo, string hoja) {
/*AGREGAR_ORIGEN*/
    }
    static void Consolidar(DataTable dtClientesCrudo, DataTable dtProductosCrudo, string ubicacionExistencias,
                           out DataTable dtClientes, out DataTable dtProductos, out DataTable dtExistencias,
                           out DataTable dtRechazos, out string resumen) {
/*CONSOLIDAR*/
    }
    // Lee una hoja exportada: valores con prefijo de tipo (n: número, b: booleano, s: texto), como los devuelve Excel.
    static DataTable Leer(string ruta) {
        var lineas = File.ReadAllLines(ruta, Encoding.UTF8);
        var t = new DataTable();
        foreach (var h in lineas[0].Split('\t')) t.Columns.Add(h, typeof(object));
        foreach (var l in lineas.Skip(1)) {
            var celdas = l.Split('\t'); var fila = t.NewRow();
            for (int i = 0; i < t.Columns.Count; i++) {
                string c = i < celdas.Length ? celdas[i] : "";
                if (c == "") fila[i] = DBNull.Value;
                else if (c.StartsWith("n:")) fila[i] = double.Parse(c.Substring(2), System.Globalization.CultureInfo.InvariantCulture);
                else if (c.StartsWith("b:")) fila[i] = c.Substring(2) == "1";
                else fila[i] = c.Substring(2).Replace("\\n", "\n").Replace("\\t", "\t");
            }
            t.Rows.Add(fila);
        }
        return t;
    }
    static void Escribir(DataTable t, string ruta) {
        var sb = new StringBuilder();
        sb.AppendLine(string.Join("\t", t.Columns.Cast<DataColumn>().Select(c => c.ColumnName).ToArray()));
        foreach (DataRow r in t.Rows)
            sb.AppendLine(string.Join("\t", r.ItemArray.Select(v => Convert.ToString(v).Replace("\n", " ")).ToArray()));
        File.WriteAllText(ruta, sb.ToString(), Encoding.UTF8);
    }
    static void Main(string[] args) {
        DataTable clientes = null, productos = null;
        foreach (var l in File.ReadAllLines(args[0], Encoding.UTF8)) {
            var p = l.Split('\t');  // hoja, archivo, ruta del .tsv
            var t = Leer(p[2]);
            AgregarOrigen(t, p[1], p[0]);
            if (p[0] == "clientes") { if (clientes == null) clientes = t; else clientes.Merge(t, false, MissingSchemaAction.Add); }
            else { if (productos == null) productos = t; else productos.Merge(t, false, MissingSchemaAction.Add); }
        }
        DataTable c, pr, e, r; string resumen;
        Consolidar(clientes, productos, args[2], out c, out pr, out e, out r, out resumen);
        Escribir(c, Path.Combine(args[1], "clientes.tsv")); Escribir(pr, Path.Combine(args[1], "productos.tsv"));
        Escribir(e, Path.Combine(args[1], "existencias.tsv")); Escribir(r, Path.Combine(args[1], "rechazos.tsv"));
        Console.OutputEncoding = Encoding.UTF8;
        Console.WriteLine(resumen);
    }
}
"""


def celda(v):
    if v is None or v == "":
        return ""
    if isinstance(v, bool):
        return "b:" + ("1" if v else "0")
    if isinstance(v, (int, float)):
        return "n:" + repr(float(v))
    return "s:" + str(v).replace("\n", "\\n").replace("\t", "\\t")


def hojas_del_libro(ruta):
    """{nombre de hoja: filas} solo para "clientes" y "productos", con valores sin formato."""
    resultado = {}
    if ruta.suffix.lower() == ".xls":
        libro = xlrd.open_workbook(ruta)
        for hoja in libro.sheets():
            if hoja.name in HOJAS:
                filas = []
                for r in range(hoja.nrows):
                    fila = []
                    for c in range(hoja.ncols):
                        x = hoja.cell(r, c)
                        fila.append(bool(x.value) if x.ctype == xlrd.XL_CELL_BOOLEAN else x.value)
                    filas.append(fila)
                resultado[hoja.name] = filas
    else:
        libro = openpyxl.load_workbook(ruta, data_only=True)
        for hoja in libro.worksheets:
            if hoja.title in HOJAS:
                resultado[hoja.title] = [list(f) for f in hoja.iter_rows(values_only=True)]
    return resultado


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    raiz = Path(sys.argv[1])
    salida = Path(sys.argv[2]) if len(sys.argv) > 2 else AQUI.parent / "salida" / "prueba"
    salida.mkdir(parents=True, exist_ok=True)
    temporal = Path(tempfile.mkdtemp(prefix="rpa_"))

    manifiesto, revisados, usados = [], 0, 0
    for ruta in sorted(raiz.rglob("*")):
        if not ruta.is_file() or ruta.suffix.lower() not in (".xls", ".xlsx") or ruta.name.startswith("~$"):
            continue
        revisados += 1
        hojas = hojas_del_libro(ruta)
        usados += bool(hojas)
        for nombre, filas in hojas.items():
            if not filas:
                continue
            tsv = temporal / f"{len(manifiesto):03d}.tsv"
            encabezado = [str(h) if h is not None else f"Columna{i}" for i, h in enumerate(filas[0])]
            lineas = ["\t".join(encabezado)] + ["\t".join(celda(v) for v in f) for f in filas[1:]]
            tsv.write_text("\n".join(lineas), encoding="utf-8")
            manifiesto.append(f"{nombre}\t{ruta}\t{tsv}")
            print(f"  {nombre:9} <- {ruta.relative_to(raiz)}")
    (temporal / "manifiesto.txt").write_text("\n".join(manifiesto), encoding="utf-8")

    codigo = PLANTILLA.replace("/*AGREGAR_ORIGEN*/", (AQUI / "AgregarOrigen.cs").read_text(encoding="utf-8"))
    codigo = codigo.replace("/*CONSOLIDAR*/", (AQUI / "Consolidar.cs").read_text(encoding="utf-8"))
    fuente, exe = temporal / "prueba.cs", temporal / "prueba.exe"
    fuente.write_text(codigo, encoding="utf-8")
    compilado = subprocess.run([str(CSC), "/nologo", f"/out:{exe}", "/r:System.Data.dll", "/r:System.Core.dll",
                                "/r:System.Xml.dll", str(fuente)], capture_output=True, text=True)
    if compilado.returncode:
        raise SystemExit("No compila el C#:\n" + compilado.stdout)
    ejecucion = subprocess.run([str(exe), str(temporal / "manifiesto.txt"), str(salida), UBICACION],
                               capture_output=True, text=True, encoding="utf-8")
    if ejecucion.returncode:
        raise SystemExit("Falló el C#:\n" + ejecucion.stdout + ejecucion.stderr)
    print(f"Archivos Excel revisados: {revisados}, con hojas útiles: {usados}")
    print(ejecucion.stdout.strip())
    print(f"Resultados en {salida}")


if __name__ == "__main__":
    main()
